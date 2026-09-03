from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import RGBColor, Pt
from docx.oxml.shared import OxmlElement


def create_table():
    # 创建新文档
    doc = Document()

    # 添加8x8表格
    table = doc.add_table(rows=8, cols=8)

    # 设置表格样式
    for row in table.rows:
        for cell in row.cells:
            # 设置单元格文本
            cell.text = '1314'

            # 获取第一个段落设置水平居中
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 设置垂直居中
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # 设置黑色边框
            set_cell_border(cell)

    # 保存文档
    doc.save('table_8x8.docx')


def set_cell_border(cell):
    # 为单元格设置四周边框
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    # 创建边框元素
    borders = ['left', 'right', 'top', 'bottom']
    for border in borders:
        tag = f'w:{border}'
        element = tcPr.find(tag)
        if element is None:
            element = OxmlElement(tag)
            tcPr.append(element)

        # 设置边框属性
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), '000000')


def qn(tag):
    # 处理XML命名空间
    return f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{tag}'


if __name__ == '__main__':
    create_table()